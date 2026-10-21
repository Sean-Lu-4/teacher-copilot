"""Printable worksheets built only from the current saved question records."""
from io import BytesIO

MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def build_docx(plan, records, audience, language='bilingual'):
    if audience not in {'student', 'teacher'}:
        raise ValueError('请选择学生版或教师版。')
    try:
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise ValueError('Word 导出需要安装项目 requirements.txt 中的依赖。') from exc
    from .worksheets import TOPICS

    document = Document()
    # The bundled default template may carry a decorative title border.
    for style in document.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in [('Normal', 10.5), ('Title', 22), ('Heading 1', 15), ('Heading 2', 11)]:
        style = document.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.paragraph_format.space_after = Pt(6)
    document.styles['Normal'].paragraph_format.line_spacing = 1.15
    if audience == 'teacher':
        # The answer key contains several explanatory blocks after each
        # question. Keep the teacher edition compact enough that a short
        # final block does not spill onto an otherwise empty extra page.
        document.styles['Normal'].font.size = Pt(9.5)
        document.styles['Normal'].paragraph_format.line_spacing = 1.0
        for name in ['Normal', 'Heading 1', 'Heading 2']:
            document.styles[name].paragraph_format.space_after = Pt(3)
        document.styles['Heading 2'].paragraph_format.space_before = Pt(4)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('AP CSA  |  ' + ('学生版' if audience == 'student' else '教师版') + '  |  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    def paragraph(text):
        return document.add_paragraph(str(text))

    def code(text):
        # One paragraph per source line permits long code to cross pages safely.
        for line in str(text).split('\n'):
            p = paragraph(line)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0 if audience == 'teacher' else 1.05
            p.paragraph_format.left_indent = Cm(.2)
            for run in p.runs:
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'F2F4F6')
            p._p.get_or_add_pPr().append(shade)
        paragraph('').paragraph_format.space_after = Pt(0)

    settings = plan.get('worksheet_settings')
    is_exam = (settings or {}).get('purpose') == 'exam'
    topic = TOPICS.get((settings or {}).get('topic'), plan.get('course', {}).get('topic', '练习'))
    document.add_paragraph('AP CSA 正式试卷' if is_exam else 'AP CSA 作业', 'Title')
    paragraph(topic + '  ' + ('学生版' if audience == 'student' else '教师答案与解析'))
    if is_exam:
        marks = sum(r['content'].get('marks', 0) for r in records)
        paragraph(f'考试时间 {settings["minutes"]} 分钟  |  总分 {marks} 分')
    if audience == 'student':
        paragraph('姓名：________________    班级：____________    日期：____________')
    for index, record in enumerate(records, 1):
        q = record['content']
        block_start = len(document.paragraphs)
        labels = {'pending': '待审核', 'passed': '通过', 'failed': '需修改'}
        heading = document.add_paragraph(f'第 {index} 题  {q["title"]}', 'Heading 1')
        if index > 1 and (q.get('question_kind') == 'write' or records[index-2]['content'].get('question_kind') == 'write'):
            heading.paragraph_format.page_break_before = True
        paragraph((f'{q.get("marks", 0)} 分  |  ' if is_exam else '') + f'预计 {q["minutes"]} 分钟')
        if not q['code'].startswith('// AP source question'):
            code(q['code'])
        if q.get('question_images'):
            from pathlib import Path
            import re
            for filename in q['question_images']:
                if not re.fullmatch(r'FRQ-[a-f0-9]{16}-S[0-9]+-p[0-9]+\.png',filename):
                    raise ValueError('Invalid question image')
                path=Path(__file__).parent/'imported_bank'/'images'/filename
                from docx.image.image import Image
                picture=Image.from_file(str(path))
                width=min(17,22*picture.px_width/picture.px_height)
                document.add_picture(str(path),width=Cm(width))
        else:
            paragraph(q['task'])
        for key, value in q.get('options', {}).items():
            paragraph(f'{key}. {value}')
        if not q.get('question_images'):
            for prompt in document.paragraphs[block_start:-1]:
                prompt.paragraph_format.keep_with_next = True
        if audience == 'student':
            document.add_paragraph('作答区', 'Heading 2')
            count = 12 if q.get('question_kind') == 'write' else 1
            for _ in range(count):
                p = paragraph('____________________________________________________________________')
                p.paragraph_format.space_after = Pt(7)
                p.runs[0].font.color.rgb = RGBColor(180, 180, 180)
            # Keep the prompt with the start of the answer area; long prompts
            # and remaining answer lines must still be able to span pages.
            for p in (document.paragraphs[block_start:-count] if not q.get('question_images') else []):
                p.paragraph_format.keep_with_next = True
            continue
        for title, value in [('参考答案', q['answer']), ('正确选项', q.get('answer_key')),
                             ('解析步骤', q['reasoning'])]:
            if language == 'english' and q.get('explanation_en'):
                continue
            if value:
                document.add_paragraph(title, 'Heading 2')
                if isinstance(value, list):
                    for n, item in enumerate(value, 1):
                        paragraph(f'{n}. {item}')
                else:
                    paragraph(value)
        if language not in {'bilingual','english','chinese'}:
            raise ValueError('解析语言选项无效。')
        if language in {'bilingual','chinese'}:
            pass
        if q.get('explanation_en') and language in {'bilingual','english'}:
            document.add_paragraph('English explanation', 'Heading 2')
            paragraph(q['explanation_en'])
        if q.get('reference_solution'):
            document.add_paragraph('参考实现', 'Heading 2')
            code(q['reference_solution'])
        for title, value in [('干扰项解释', [f'{k}：{v}' for k, v in q.get('distractors', {}).items()]),
                             ('评分标准', q['rubric'])]:
            if value:
                document.add_paragraph(title, 'Heading 2')
                for item in value if isinstance(value, list) else [value]:
                    paragraph(item)
    document.core_properties.title = 'AP CSA ' + topic + ' ' + audience
    document.core_properties.author = 'CSA Studio'
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
