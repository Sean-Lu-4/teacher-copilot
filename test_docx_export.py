from test_pdf_fixtures import pdf_ids
from teacher_copilot.worksheets import TOPIC_CONTENTS
import base64
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from docx import Document
from teacher_copilot.review_app import ReviewStore, Conflict
from teacher_copilot.worksheets import TOPIC_CONTENTS


class WordExportTests(unittest.TestCase):
    def test_selected_exam_edit_save_and_export(self):
        with tempfile.TemporaryDirectory() as root:
            store = ReviewStore(root)
            state = store.mutate(dict(action='generate', settings=dict(purpose='exam',
                topic='csa-calls', contents=TOPIC_CONTENTS['csa-calls'],
                difficulty='mixed', minutes=45, question_count=3, total_marks=20,
                question_types=['concept','write'], selected_ids=pdf_ids())))
            q = state['records'][0]['content']
            state = store.mutate(dict(action='revise', run=state['run'], version=state['version'],
                id=q['id'], actor='Test teacher', reason='Exam export acceptance',
                changes=dict(title='Exam saved edit', answer='EXAM_PRIVATE_ANSWER')))
            self.assertEqual(sum(r['content']['marks'] for r in state['records']),20)
            for audience in ['student','teacher']:
                result=store.mutate(dict(action='export', format='docx', audience=audience,
                    run=state['run'], version=state['version']))
                doc=Document(BytesIO(base64.b64decode(result['download'])))
                text='\n'.join(p.text for p in doc.paragraphs)
                self.assertIn('Exam saved edit',text)
                self.assertIn('总分 20 分',text)
                self.assertEqual('EXAM_PRIVATE_ANSWER' in text,audience=='teacher')

    def test_exam_marks_and_time_match_saved_questions(self):
        from teacher_copilot.worksheets import assemble_exam
        from teacher_copilot.docx_export import build_docx
        settings, records = assemble_exam(dict(purpose='exam', topic='csa-calls',
            contents=TOPIC_CONTENTS['csa-calls'], difficulty='mixed',
            minutes=60, question_count=10, total_marks=40,
            question_types=['write', 'concept']))
        for audience in ['student', 'teacher']:
            doc = Document(BytesIO(build_docx({'worksheet_settings': settings}, records, audience)))
            text = '\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('AP CSA 正式试卷', text)
            self.assertIn('考试时间 60 分钟  |  总分 40 分', text)
            for record in records:
                q = record['content']
                self.assertIn(f'{q["marks"]} 分  |  预计 {q["minutes"]} 分钟', text)
            headings = [p for p in doc.paragraphs if p.style.name == 'Heading 1']
            for i, heading in enumerate(headings[1:], 1):
                expected = any(records[j]['content'].get('question_kind') == 'write' for j in [i-1, i])
                self.assertEqual(bool(heading.paragraph_format.page_break_before), expected)

    def test_latest_saved_content_and_student_answer_isolation(self):
        with tempfile.TemporaryDirectory() as root:
            store = ReviewStore(root)
            state = store.mutate(dict(action='generate', settings=dict(topic='csa-arrays',
                contents=TOPIC_CONTENTS['csa-arrays'], difficulty='any', minutes=20, purpose='homework')))
            old = state
            question = state['records'][0]['content']
            state = store.mutate(dict(action='revise', run=state['run'], version=state['version'],
                id=question['id'], actor='Test teacher', reason='Verify edited Word exports', changes=dict(title='SAVED_TITLE', answer='PRIVATE_ANSWER',
                rubric='PRIVATE_RUBRIC', reasoning=['PRIVATE_REASONING'], intervention='PRIVATE_GUIDANCE')))
            for audience in ['student', 'teacher']:
                result = store.mutate(dict(action='export', format='docx', audience=audience,
                    run=state['run'], version=state['version']))
                binary = base64.b64decode(result['download'])
                self.assertEqual(binary, (Path(result['export_path']) / result['filename']).read_bytes())
                text = '\n'.join(p.text for p in Document(BytesIO(binary)).paragraphs)
                self.assertIn('SAVED_TITLE', text)
                self.assertNotIn('// AP source question', text)
                for secret in ['PRIVATE_ANSWER', 'PRIVATE_RUBRIC', 'PRIVATE_REASONING']:
                    self.assertEqual(secret in text, audience == 'teacher')
            with self.assertRaises(Conflict):
                store.mutate(dict(action='export', format='docx', audience='teacher',
                    run=old['run'], version=old['version']))
            with self.assertRaises(ValueError):
                store.mutate(dict(action='export', format='pdf', audience='teacher',
                    run=state['run'], version=state['version']))


if __name__ == '__main__':
    unittest.main()
